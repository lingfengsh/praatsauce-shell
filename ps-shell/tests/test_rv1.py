# -*- coding: utf-8 -*-
"""PS-RV1 审核结果复用（读对象，不重算 F0/共振峰）—— 单元测试。

覆盖本轮的三件事：
1. 审查抽取参数（F0 上下限 / 共振峰上限 / 数量）由调用方传入，不再写死；
2. 导出对象的时间步 = 批量的 intervalFixed（否则读回时帧数不足会被补 0）；
3. 接口层把这些参数正确透传（key 名写错就会静默退回默认值，最容易出漏洞）。

**不依赖 Praat**：所有会落到 Praat 的调用都换成打桩（`praat.run_script` /
`review_export._run_batch`），只验证"传了什么参数、写了什么清单、产物怎么判定"。
也不读写真实用户目录：全部在 tempfile 里。

运行：
    cd ps-shell
    python -m unittest tests.test_rv1 -v
"""

import os
import re
import sys
import json
import shutil
import tempfile
import threading
import unittest
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import review_grid, review_export, server  # noqa: E402
from backend import vowel_ref, review_choice, formant_path  # noqa: E402
from backend import series_export, praat_objects  # noqa: E402

def _mp(d, name):
    """measure/ 子目录里的夹具路径（2026-09-27 方案 A 布局：对象文件归子目录）。"""
    md = os.path.join(d, "measure")
    os.makedirs(md, exist_ok=True)
    return os.path.join(md, name)




class ScriptParamsTest(unittest.TestCase):
    """生成的 Praat 脚本必须用传进来的参数，而不是写死的 75/500/5500/5。"""

    ITEMS = [(os.path.join("X:", "d", "11.wav"), "11")]

    # 基频抽取统一走滤波自相关（2026-09-25 用户拍板，与主批量 pitchMethod=0 一致）
    FILTERED = "To Pitch (filtered autocorrelation): %s, %g, %g, 15, 0, 0.03, 0.09, 0.5, 0.055, 0.35, 0.14"

    def test_defaults_unchanged(self):
        """不传参数时与改动前完全一致（默认 75 / 500 / 5500 / 5）。"""
        s = review_grid._build_script(self.ITEMS, "OUT", 500.0, 5500.0)
        self.assertIn(self.FILTERED % ("0.0", 75.0, 500.0), s)
        self.assertIn("To Formant (burg): 0.0, 5, 5500, 0.025, 50", s)

    def test_custom_params_injected(self):
        s = review_grid._build_script(self.ITEMS, "OUT", 300.0, 5000.0,
                                      pitch_floor=50.0, n_formants=5.0)
        self.assertIn(self.FILTERED % ("0.0", 50.0, 300.0), s)
        self.assertIn("To Formant (burg): 0.0, 5, 5000, 0.025, 50", s)
        self.assertNotIn("75", s.split("To Pitch")[1].split("\n")[0])

    def test_half_integer_n_formants(self):
        """Praat 手册允许数量取 0.5 的倍数，不能被 %d 截断。"""
        s = review_grid._build_script(self.ITEMS, "OUT", 500.0, 5500.0, n_formants=5.5)
        self.assertIn("To Formant (burg): 0.0, 5.5, 5500, 0.025, 50", s)

    def test_nf_script_params(self):
        """「多数量择优」：每个文件按 4/5/6/7 各跑一套 To Formant (burg)，Pitch 只抽一次。"""
        s = review_grid._build_nf_script(self.ITEMS, "OUT", 300.0, 5000.0, pitch_floor=50.0)
        self.assertIn(self.FILTERED % ("0.0", 50.0, 300.0), s)
        for nf in review_grid.NF_CANDIDATES:
            self.assertIn("To Formant (burg): 0.0, %d, 5000, 0.025, 50" % nf, s)
        # Pitch 每文件只抽一次；共振峰每个数量各一份 .F<nf>.formant
        self.assertEqual(s.count(".pitch"), len(self.ITEMS))
        self.assertEqual(s.count(".F4.formant"), len(self.ITEMS))
        self.assertEqual(s.count(".F7.formant"), len(self.ITEMS))

    def test_track_snippet_still_attached(self):
        """加了参数之后，Track 片段仍要接在 To Formant 之后。"""
        s = review_grid._build_script(self.ITEMS, "OUT", 500.0, 5500.0,
                                      formant_track=True, track_refs={"11": [0.1, 0.5]})
        self.assertIn("To Formant (burg)", s)
        self.assertIn("Track", s)

    def test_track_uses_ipa_refs(self):
        """有音标 → Track 参考频率用音标期望值（字面赋值）；没音标 → Get mean 自动估计。"""
        s = review_grid._build_script(self.ITEMS, "OUT", 500.0, 5500.0,
                                      formant_track=True, track_refs={"11": [0.1, 0.5]},
                                      track_ipa={"11": (300.0, 900.0, 2200.0)})
        self.assertIn("rf1 = 300", s)
        self.assertIn("rf2 = 900", s)
        self.assertIn("rf3 = 2200", s)
        self.assertNotIn("Get mean: 1", s)
        s2 = review_grid._build_script(self.ITEMS, "OUT", 500.0, 5500.0,
                                       formant_track=True, track_refs={"11": [0.1, 0.5]})
        self.assertIn('Get mean: 1, 0.1, 0.5, "Hertz"', s2)

    def test_ipa_refs_helper(self):
        """_ipa_refs：取首段元音期望；无元音 / 库外音标 → None（回退自动估计）。"""
        r_i = review_grid._ipa_refs("i", "male")
        self.assertIsNotNone(r_i)
        self.assertEqual(r_i, review_grid._ipa_refs(["i"], "male"))  # 字符串/列表同解析
        # 双元音取首段（Track 一次只吃一套参考值，与 review_build 推荐组同约定）
        self.assertEqual(review_grid._ipa_refs("iu", "male"), r_i)
        self.assertTrue(all(v is None or v > 0 for v in r_i))
        self.assertIsNone(review_grid._ipa_refs("xk", None))   # 无已知元音字符
        self.assertIsNone(review_grid._ipa_refs("", None))


class GridEchoTest(unittest.TestCase):
    """build_grid 要把本次实际用的参数回显——前端靠它设语图纵轴。"""

    def test_empty_dir_echoes_params(self):
        with tempfile.TemporaryDirectory() as d:
            r = review_grid.build_grid(d, pitch_floor=50.0, pitch_ceiling=300.0,
                                       formant_ceiling=5000.0, n_formants=5.0)
        self.assertTrue(r["ok"])
        self.assertEqual(r["total"], 0)
        self.assertEqual(r["pitchFloor"], 50.0)
        self.assertEqual(r["pitchCeiling"], 300.0)
        self.assertEqual(r["formantCeiling"], 5000.0)
        self.assertEqual(r["nFormants"], 5.0)

    def test_missing_dir_is_error(self):
        r = review_grid.build_grid(os.path.join("X:", "no", "such", "dir"))
        self.assertFalse(r["ok"])
        self.assertIn("目录不存在", r["error"])


class GridReuseTest(unittest.TestCase):
    """「有存盘中间文件优先复用」（2026-09-26 用户拍板）：

    - wav 旁 .Pitch/.Formant 都在 → 不跑 Praat，直接读盘展示；
    - 缺哪维补哪维 → 现算的结果复制到 wav 旁存盘；已有的文件一个字节都不动。
    用假 _run_script 模拟 Praat（家规：tests/ 内不跑真实 Praat）。
    """

    PITCH_TEXT = (
        'File type = "ooTextFile"\n'
        'Object class = "Pitch 1"\n\n'
        'xmin = 0\nxmax = 0.1\nnx = 3\ndx = 0.01\n'
        'x1 = 0.020000000000000014\nceiling = 500\nmaxnCandidates = 15\n'
        'frames []:\n'
        '    frames [1]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 100\n            strength = 0.9\n'
        '    frames [2]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 120\n            strength = 0.9\n'
        '    frames [3]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 110\n            strength = 0.9\n'
    )

    def setUp(self):
        # 假 Praat：记录调用次数与脚本文本，并把"现算结果"写进临时输出目录
        self.calls = []
        self.scripts = []
        self._orig = (review_grid._resolve_exe, review_grid._run_script,
                      review_grid.audio._wav_info,
                      review_grid.praat_objects.read_formant_text)
        review_grid._resolve_exe = lambda exe=None: "fake-prat.exe"

        def fake_run(exe, out_dir, script, name):
            self.calls.append(name)
            self.scripts.append(script)
            with open(os.path.join(out_dir, "a.pitch"), "w", encoding="utf-8") as f:
                f.write(self.PITCH_TEXT.replace("frequency = 100", "frequency = 300"))
            with open(os.path.join(out_dir, "a.formant"), "w", encoding="utf-8") as f:
                # 平滑升级脚本产出"已 Track"的对象（真机上是 Track 后的 3 条轨迹）
                f.write("TRACKED" if name == "grid_upgrade.praat" else "")
            return None, "OK grid done"

        review_grid._run_script = fake_run
        review_grid.audio._wav_info = lambda p: (0.5, 1, 16000)
        review_grid.praat_objects.read_formant_text = \
            lambda p: [{"t": 0.01, "F1": 500.0, "F2": 1500.0, "F3": 2500.0}]

    def tearDown(self):
        (review_grid._resolve_exe, review_grid._run_script,
         review_grid.audio._wav_info,
         review_grid.praat_objects.read_formant_text) = self._orig

    def _script(self):
        return "\n".join(self.scripts)

    def _make_dir(self, with_pitch=False, with_formant=False):
        d = tempfile.mkdtemp(prefix="reuse_")
        open(os.path.join(d, "a.wav"), "wb").close()
        if with_pitch:
            with open(_mp(d, "a.Pitch"), "w", encoding="utf-8") as f:
                f.write(self.PITCH_TEXT)
        if with_formant:
            open(_mp(d, "a.Formant"), "w").close()
        return d

    def test_both_on_disk_skip_praat(self):
        d = self._make_dir(with_pitch=True, with_formant=True)
        try:
            r = review_grid.build_grid(d)
            self.assertTrue(r["ok"])
            # 不再是"一个 Praat 都不跑"（2026-09-26 平滑升级拍板）：盘上旧裸 burg 对象
            # 无审定标记 → 补跑一次 Track 升级并写回，之后才真正零 Praat。
            self.assertEqual(self.calls, ["grid_upgrade.praat"])
            self.assertEqual(r["computed"], 0)
            self.assertEqual(r["formantUpgraded"], 1)
            self.assertEqual(r["pitchReused"], 1)
            f0 = r["files"][0]
            self.assertTrue(f0["pitchSaved"] and f0["formantSaved"])
            self.assertEqual([p["f0"] for p in f0["pitch"]], [100.0, 120.0, 110.0])
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_upgrade_writes_smoothed_object(self):
        """复用对象平滑升级（2026-09-26 用户反馈 12.wav 九宫格不平滑、详情却平滑）：
        无审定标记的盘上 .Formant 补跑 Track，平滑对象 copy2 写回，下次零开销。"""
        d = self._make_dir(with_pitch=True, with_formant=True)
        try:
            with open(_mp(d, "a.Formant"), "w", encoding="utf-8") as f:
                f.write("RAW-OLD-OBJECT")
            r = review_grid.build_grid(d)
            self.assertTrue(r["ok"])
            self.assertEqual(r["formantUpgraded"], 1)
            with open(_mp(d, "a.Formant"), encoding="utf-8") as f:
                self.assertEqual(f.read(), "TRACKED")
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_upgrade_respects_audited_choices(self):
        """已审定文件不强加平滑：
        - formantSmoothed=false（审定明确不要平滑）→ 完全不跑 Praat；
        - formant 已审定（无论标记）→ 不升级；
        - formantSmoothed=true（盘上已是平滑版）→ 不升级。"""
        d = self._make_dir(with_pitch=True, with_formant=True)
        try:
            cases = [
                {"pitch": {"id": "500", "label": "x"}, "formant": None,
                 "formantSmoothed": False},
                {"formant": {"id": "5500-5", "label": "x"}, "formantSmoothed": False},
                {"formant": {"id": "5500-5", "label": "x"}, "formantSmoothed": True},
                {"formant": {"id": "5500-5", "label": "x"}},
            ]
            for ch in cases:
                self.calls.clear()
                with open(_mp(d, "a.measure.json"), "w",
                          encoding="utf-8") as f:
                    json.dump(ch, f)
                r = review_grid.build_grid(d)
                self.assertTrue(r["ok"])
                self.assertEqual(self.calls, [], "审定文件不应触发升级：%s" % ch)
                self.assertEqual(r["formantUpgraded"], 0)
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_missing_formant_computed_and_saved(self):
        d = self._make_dir(with_pitch=True)
        try:
            r = review_grid.build_grid(d)
            self.assertTrue(r["ok"])
            self.assertEqual(len(self.calls), 1)      # 只为缺的共振峰跑一次
            self.assertEqual(r["computed"], 1)
            # 只补缺的一维：.Pitch 已在盘 → 脚本里不得再有 Pitch 抽取调用
            # （2026-09-26 用户批评：重算已存盘的一维纯属浪费）
            self.assertNotIn("To Pitch", self._script())
            # 现算的 .Formant 已存盘；已有 .Pitch 一字未动（不是现算值 300）
            self.assertTrue(os.path.isfile(_mp(d, "a.Formant")))
            with open(_mp(d, "a.Pitch"), encoding="utf-8") as f:
                self.assertIn("frequency = 100", f.read())
            f0 = r["files"][0]
            self.assertTrue(f0["pitchSaved"] and f0["formantSaved"])
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_reuse_false_recomputes(self):
        d = self._make_dir(with_pitch=True, with_formant=True)
        try:
            r = review_grid.build_grid(d, reuse=False)
            self.assertEqual(len(self.calls), 1)
            self.assertEqual(r["computed"], 1)
            self.assertFalse(r["files"][0]["pitchSaved"])
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_reuse_formant_opt_path(self):
        """复用 + 多数量择优：need_tgs 必须保持 (tg_path, intervals) 元组形状
        （2026-09-26 踩坑：写成裸路径字符串，_merge_nf 解包报
        "too many values to unpack (expected 2)"）。mock 下没有 .F*.formant，
        合并走 failed 分支——但解包行会被真实执行，形状错了立刻炸。"""
        d = self._make_dir()
        try:
            r = review_grid.build_grid(d, formant_opt=True)
            self.assertTrue(r["ok"])
            self.assertEqual(len(self.calls), 1)
            self.assertEqual(r["files"][0]["formant"], [])   # 无候选文件 → 合并为空
        finally:
            shutil.rmtree(d, ignore_errors=True)


class ExportDefaultsTest(unittest.TestCase):
    """导出参数的解析：默认值、覆盖、坏值回退。"""

    def test_params_for_defaults(self):
        self.assertEqual(review_export.params_for(None), (500.0, 5, 5500.0))

    def test_params_for_choice(self):
        self.assertEqual(
            review_export.params_for({"pitch": {"id": "300"}, "formant": {"id": "5000-6"}}),
            (300.0, 6, 5000.0))

    def test_params_for_legacy_bare_ceiling(self):
        """旧记录只存了裸上限（没有 -数量），数量退回 5。"""
        self.assertEqual(review_export.params_for({"formant": {"id": "5000"}}),
                         (500.0, 5, 5000.0))

    def test_pitch_floor_default_and_override(self):
        self.assertEqual(review_export.pitch_floor_for(None), 75.0)
        self.assertEqual(review_export.pitch_floor_for({"pitchFloor": 50}), 50.0)
        self.assertEqual(review_export.pitch_floor_for({"pitchFloor": "bad"}), 75.0)
        self.assertEqual(review_export.pitch_floor_for({"pitchFloor": 0}), 75.0)

    def test_time_step_default_and_override(self):
        self.assertEqual(review_export.time_step_for(None), 0.005)
        self.assertEqual(review_export.time_step_for({"timeStep": 0.01}), 0.01)
        self.assertEqual(review_export.time_step_for({"timeStep": 0}), 0.005)
        self.assertEqual(review_export.time_step_for({"timeStep": None}), 0.005)


class RunBatchArgsTest(unittest.TestCase):
    """跑 Praat 时传下去的参数与清单内容（这是最容易传错的一层）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.captured = {}
        self._orig = review_export.praat.run_script

        def fake_run(exe, script, args, cwd=None, full_trust=False, timeout=None):
            self.captured["script"] = script
            self.captured["args"] = list(args)
            wl = args[1]
            if os.path.isfile(wl):
                with open(wl, encoding="utf-8") as f:
                    self.captured["worklist"] = f.read()
            else:
                self.captured["worklist"] = ""
            return 0, [{"text": "EXPORT_DONE"}]

        review_export.praat.run_script = fake_run

    def tearDown(self):
        review_export.praat.run_script = self._orig
        self.tmp.cleanup()

    def test_export_one_passes_floor_and_step(self):
        wav = os.path.join(self.tmp.name, "11.wav")
        open(wav, "wb").close()
        r = review_export.export_one(wav, None,
                                    {"pitchFloor": 50, "timeStep": 0.01,
                                     "formantCeiling": 5000, "nFormants": 5}, exe="P")
        args = self.captured["args"]
        self.assertEqual(args[3], "50")          # f0min
        self.assertEqual(args[4], "0.01")        # timeStep
        # 2026-09-25 起清单固定 11 段：后 7 段是自定义基频参数，普通候选全空
        self.assertEqual(self.captured["worklist"], "11.wav|500|5|5000|||||||\n")
        self.assertEqual(r["params"]["pitchFloor"], 50.0)
        self.assertEqual(r["params"]["timeStep"], 0.01)

    def test_pitch_params_for(self):
        """自定义基频参数（详情页滑块存盘）的解析：合法/缺项/普通候选。"""
        self.assertIsNone(review_export.pitch_params_for(None))
        self.assertIsNone(review_export.pitch_params_for({"pitch": {"id": "500"}}))
        self.assertIsNone(review_export.pitch_params_for(
            {"pitch": {"id": "custom", "method": "raw"}}))          # 缺 params
        self.assertIsNone(review_export.pitch_params_for(
            {"pitch": {"id": "custom", "method": "bad",
                       "params": {"method": "bad", "ceiling": 600}}}))  # 非法 method
        ch = {"pitch": {"id": "custom", "method": "filtered",
                        "params": {"method": "filtered", "ceiling": 800, "silence": 0.09,
                                   "voicing": 0.5, "octave": 0.055, "octaveJump": 0.35,
                                   "vuv": 0.14, "attenuation": 0.03}}}
        out = review_export.pitch_params_for(ch)
        self.assertEqual(out["method"], "filtered")
        self.assertEqual(out["ceiling"], 800.0)
        self.assertEqual(out["silence"], 0.09)

    def test_export_one_custom_pitch_params_worklist(self):
        """滑块自定义参数存盘后，导出清单应带上 method 与全部 7 个参数。"""
        wav = os.path.join(self.tmp.name, "12.wav")
        open(wav, "wb").close()
        choice = {"pitch": {"id": "custom", "label": "自定义参数（raw AC·上限600Hz）",
                            "method": "raw",
                            "params": {"method": "raw", "ceiling": 600, "silence": 0.05,
                                       "voicing": 0.4, "octave": 0.02, "octaveJump": 0.4,
                                       "vuv": 0.1, "attenuation": 0.03}}}
        r = review_export.export_one(wav, choice, None, exe="P")
        self.assertIn("12.wav|600|5|5500|raw|0.05|0.4|0.02|0.4|0.1|0.03",
                      self.captured["worklist"])
        self.assertEqual(r["params"]["pitchCeiling"], 600.0)   # pc 被自定义 ceiling 覆盖

    def test_export_one_dims_in_worklist(self):
        """2026-09-25 存盘分维：dims=pitch/formant 写进清单第 12 段；both 不写。"""
        wav = os.path.join(self.tmp.name, "13.wav")
        open(wav, "wb").close()
        review_export.export_one(wav, None, None, exe="P", dims="pitch")
        self.assertEqual(self.captured["worklist"], "13.wav|500|5|5500||||||||pitch\n")
        review_export.export_one(wav, None, None, exe="P", dims="formant")
        self.assertEqual(self.captured["worklist"], "13.wav|500|5|5500||||||||formant\n")
        review_export.export_one(wav, None, None, exe="P", dims="both")
        self.assertEqual(self.captured["worklist"], "13.wav|500|5|5500|||||||\n")
        # 非法值回落 both
        review_export.export_one(wav, None, None, exe="P", dims="nonsense")
        self.assertEqual(self.captured["worklist"], "13.wav|500|5|5500|||||||\n")

    def test_export_dir_writes_one_line_per_wav(self):
        d = self.tmp.name
        for n in ("a.wav", "b.wav"):
            open(os.path.join(d, n), "wb").close()
        review_export.export_dir(d, defaults={"pitchFloor": 50}, exe="P")
        lines = [x for x in self.captured["worklist"].splitlines() if x]
        self.assertEqual(len(lines), 2)
        for ln in lines:
            # 2026-09-25 起固定 11 段：名|上限|数量|共振峰上限|method|silence|voicing|
            # octave|octaveJump|vuv|attenuation（无自定义参数时后 7 段全空）
            self.assertEqual(len(ln.split("|")), 11)
            self.assertTrue(ln.endswith("||||||"))
        self.assertEqual(self.captured["args"][3], "50")


class ReviewedStateTest(unittest.TestCase):
    """「已审核」= .Pitch 与 .Formant 都在（文件即事实）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        for n in ("a.wav", "b.wav", "c.wav"):
            open(os.path.join(self.d, n), "wb").close()

    def tearDown(self):
        self.tmp.cleanup()

    def test_none_reviewed(self):
        st = review_export.reviewed_state(self.d)
        self.assertEqual((st["total"], st["reviewed"]), (3, 0))
        self.assertEqual(sorted(st["missing"]), ["a.wav", "b.wav", "c.wav"])

    def test_partial(self):
        open(_mp(self.d, "a.Pitch"), "w").close()
        open(_mp(self.d, "a.Formant"), "w").close()
        # b 只有一半 → 不算审核过
        open(_mp(self.d, "b.Pitch"), "w").close()
        st = review_export.reviewed_state(self.d)
        self.assertEqual((st["total"], st["reviewed"]), (3, 1))
        self.assertEqual(sorted(st["missing"]), ["b.wav", "c.wav"])

    def test_recursive(self):
        sub = os.path.join(self.d, "sub")
        os.makedirs(sub)
        open(os.path.join(sub, "d.wav"), "wb").close()
        st = review_export.reviewed_state(self.d)
        self.assertEqual(st["total"], 4)
        self.assertIn("sub/d.wav", st["missing"])


class ManualPitchTest(unittest.TestCase):
    """人工修正写进 .Pitch 文本（只改最近帧的 candidates[1].frequency）。"""

    TEXT = (
        'File type = "ooTextFile"\n'
        'Object class = "Pitch 1"\n\n'
        'xmin = 0\nxmax = 0.1\nnx = 3\ndx = 0.01\n'
        'x1 = 0.020000000000000014\nceiling = 500\nmaxnCandidates = 15\n'
        'frames []:\n'
        '    frames [1]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 100\n            strength = 0.9\n'
        '    frames [2]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 0\n            strength = 0\n'
        '    frames [3]:\n        intensity = 0.5\n        nCandidates = 1\n'
        '        candidates [1]:\n            frequency = 120\n            strength = 0.9\n'
    )

    def _write(self, d):
        p = os.path.join(d, "t.Pitch")
        with open(p, "w", encoding="utf-8") as f:
            f.write(self.TEXT)
        return p

    def _read(self, p):
        with open(p, encoding="utf-8") as f:
            return f.read()

    def test_snap_to_nearest_frame(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            n = review_export.apply_manual_pitch(p, [{"t": 0.021, "f0": 310}])
            self.assertEqual(n, 2)                     # 帧1 手工值 + 帧2 无声帧桥接
            txt = self._read(p)
        self.assertIn("frequency = 310", txt)
        self.assertIn("frequency = 120", txt)          # 有声帧3 不受影响
        # 帧2（无声）落在手工点 310 与存活帧 120 之间 → 桥接填充（与前端显示一致）
        freqs = re.findall(r"frequency = ([\d.]+)", txt)
        self.assertEqual(len(freqs), 3)
        v2 = float(freqs[1])
        self.assertGreater(v2, 120.0)
        self.assertLess(v2, 310.0)

    def test_turn_unvoiced_into_voiced(self):
        """原本无声的帧被改成有声时，strength 也要给非零值。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            review_export.apply_manual_pitch(p, [{"t": 0.03, "f0": 200}])
            txt = self._read(p)
        self.assertIn("frequency = 200", txt)
        self.assertIn("strength = 0.5", txt)

    def test_out_of_range_point_ignored(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            self.assertEqual(review_export.apply_manual_pitch(p, [{"t": 99.0, "f0": 300}]), 0)

    def test_unvoiced_gap_bridged_between_two_points(self):
        """2026-09-26 第四次拍板（所见=所存）：两点之间的**无声帧**用 PCHIP 桥接值填充
        （与前端 mergedF0Points 同一算法）——存盘读回的曲线连续穿过手工点，
        九宫格不再因「孤立点连不出线」而看不见手工点。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            # 两点：t=0.02（帧1）=300、t=0.04（帧3）=100 → 中间无声帧2 = 桥接 200
            n = review_export.apply_manual_pitch(
                p, [{"t": 0.020, "f0": 300}, {"t": 0.040, "f0": 100}])
            self.assertEqual(n, 3)
            txt = self._read(p)
        self.assertIn("frequency = 300", txt)
        self.assertIn("frequency = 100", txt)
        self.assertIn("frequency = 200", txt)     # 帧2 = PCHIP 桥接值

    def test_voiced_frames_never_interpolated(self):
        """2026-09-26 拍板不变式：**有声原始测量帧**绝不被插值线抹掉（用户实测否决过
        「整段替换」）。fixture 帧2 是无声的，这里换成有声版验证原样保留。"""
        voiced = self.TEXT.replace("frequency = 0\n            strength = 0",
                                   "frequency = 110\n            strength = 0.8")
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "t.Pitch")
            with open(p, "w", encoding="utf-8") as f:
                f.write(voiced)
            review_export.apply_manual_pitch(
                p, [{"t": 0.020, "f0": 300}, {"t": 0.040, "f0": 100}])
            txt = self._read(p)
        self.assertIn("frequency = 110", txt)     # 有声帧2 原样保留
        self.assertNotIn("frequency = 200", txt)  # 没有插值线覆盖它

    def test_deleted_frame_bridged(self):
        """2026-09-26 新增：删掉的帧用相邻存活帧的 PCHIP 桥接值写回（与前端显示同一算法）。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            # 删帧2（t=0.03）：存活帧 1、3 值 100/120 → 桥接 = 110
            n = review_export.apply_manual_pitch(p, [], [0.03])
            self.assertEqual(n, 1)
            txt = self._read(p)
        self.assertIn("frequency = 110", txt)

    def test_deleted_with_manual_anchor(self):
        """删帧的桥接把手工覆盖值当作节点：帧1 手工改 300，删帧2 → 桥接 ∈ (120, 300)。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            review_export.apply_manual_pitch(
                p, [{"t": 0.020, "f0": 300}], [0.03])
            txt = self._read(p)
        freqs = re.findall(r"frequency = ([\d.]+)", txt)
        self.assertEqual(len(freqs), 3)
        self.assertEqual(float(freqs[0]), 300.0)   # 帧1 = 手工值
        v = float(freqs[1])                        # 帧2 = 桥接值
        self.assertGreater(v, 120.0)
        self.assertLess(v, 300.0)
        self.assertEqual(float(freqs[2]), 120.0)   # 帧3 不受影响

    def test_manual_wins_over_deleted_same_frame(self):
        """同一帧既删又加点 = 改值（手工点优先）。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            n = review_export.apply_manual_pitch(
                p, [{"t": 0.020, "f0": 300}], [0.02])
            self.assertEqual(n, 2)                     # 帧1 手工值 + 帧2 无声帧桥接
            txt = self._read(p)
        self.assertIn("frequency = 300", txt)
        # 帧2（无声）落在手工点 300 与存活帧 120 之间 → 桥接（2026-09-26 所见=所存）
        freqs = re.findall(r"frequency = ([\d.]+)", txt)
        v2 = float(freqs[1])
        self.assertGreater(v2, 120.0)
        self.assertLess(v2, 300.0)

    def test_points_outside_span_keep_curve(self):
        """同帧两点后者覆盖前者；有声帧不受插值影响。"""
        with tempfile.TemporaryDirectory() as d:
            p = self._write(d)
            # 两点都吸附到帧1（后者覆盖前者）→ 帧2 桥接、帧3 保持原样
            n = review_export.apply_manual_pitch(
                p, [{"t": 0.0195, "f0": 300}, {"t": 0.0205, "f0": 280}])
            self.assertEqual(n, 2)
            txt = self._read(p)
        self.assertIn("frequency = 280", txt)      # 同帧两点，后点覆盖
        self.assertIn("frequency = 120", txt)      # 有声帧3 不受影响
        self.assertNotIn("frequency = 300", txt)
        self.assertNotIn("frequency = 100", txt)   # 帧1 原 100 被手工值 280 覆盖

    def test_missing_file_is_noop(self):
        self.assertEqual(review_export.apply_manual_pitch("X:/no/such.Pitch",
                                                          [{"t": 0, "f0": 1}]), 0)


class PchipTest(unittest.TestCase):
    """PCHIP（Fritsch-Carlson 保单调三次插值）核心，前后端共用同一算法。"""

    def _derivs(self, xs, ys):
        return review_export._pchip_derivs(xs, ys)

    def _eval(self, xs, ys, t):
        return review_export._pchip_eval(xs, ys, self._derivs(xs, ys), t)

    def test_two_points_degenerate_linear(self):
        """只有两个手工点时 PCHIP 退化成线性插值。"""
        xs, ys = [0.0, 1.0], [100.0, 200.0]
        self.assertAlmostEqual(self._eval(xs, ys, 0.5), 150.0, places=6)
        self.assertAlmostEqual(self._eval(xs, ys, 0.0), 100.0, places=6)
        self.assertAlmostEqual(self._eval(xs, ys, 1.0), 200.0, places=6)

    def test_pass_through_nodes(self):
        """插值必须过所有手工点（节点值保持）。"""
        xs, ys = [0.0, 1.0, 2.0], [100.0, 200.0, 100.0]
        for x, y in zip(xs, ys):
            self.assertAlmostEqual(self._eval(xs, ys, x), y, places=6)

    def test_peak_no_overshoot(self):
        """峰值数据（100→200→100）：中间帧平滑过渡且不超过峰 200（不过冲）。"""
        xs, ys = [0.0, 1.0, 2.0], [100.0, 200.0, 100.0]
        for t in (0.3, 0.5, 0.7, 1.3, 1.5, 1.7):
            v = self._eval(xs, ys, t)
            self.assertTrue(100.0 < v < 200.0, "t=%s 插值 %s 应落在 (100,200)" % (t, v))

    def test_strict_monotone_preserved(self):
        """数据严格单调时，插值在任意中间点也保持单调不减（不过冲、不回摆）。"""
        xs = [0.0, 0.5, 1.0, 1.5, 2.0]
        ys = [10.0, 20.0, 30.0, 40.0, 50.0]
        prev = -1e9
        for t in (0.1, 0.4, 0.6, 0.9, 1.1, 1.4, 1.9):
            v = self._eval(xs, ys, t)
            self.assertGreater(v, prev, "保单调：t=%s 的插值 %s 应 > 前一值" % (t, v))
            prev = v


class HttpRoutingTest(unittest.TestCase):
    """接口层 key 名透传——名字写错不会报错，只会静默用默认值（最难发现的漏洞）。"""

    @classmethod
    def setUpClass(cls):
        cls._orig_build = review_grid.build_grid
        cls._orig_export_dir = review_export.export_dir
        cls._seen = {}

        def fake_build(dir_path, **kw):
            cls._seen["build"] = kw
            return {"ok": True, "dir": dir_path, "files": [], "total": 0, "pages": 0}

        def fake_export_dir(dir_path, defaults=None, **kw):
            cls._seen["export_dir"] = {"dir": dir_path, "defaults": defaults}
            return {"ok": True, "total": 0, "done": 0, "failed": []}

        review_grid.build_grid = fake_build
        review_export.export_dir = fake_export_dir
        cls.srv = server.start_server(0)
        cls.base = "http://127.0.0.1:%d" % cls.srv.server_address[1]

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        review_grid.build_grid = cls._orig_build
        review_export.export_dir = cls._orig_export_dir

    def _post(self, path, payload):
        req = urllib.request.Request(
            self.base + path, data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json",
                     "X-PS-Token": server.TOKEN},   # 2026-09-28 起 /api/* 要求一次性 Token
            method="POST")
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.loads(r.read().decode("utf-8"))

    def test_grid_receives_four_params(self):
        self._post("/api/review/grid", {"dir": "D", "pitchFloor": 50,
                                        "pitchCeiling": 300, "formantCeiling": 5000,
                                        "nFormants": 5})
        kw = self._seen["build"]
        self.assertEqual(kw["pitch_floor"], 50.0)
        self.assertEqual(kw["pitch_ceiling"], 300.0)
        self.assertEqual(kw["formant_ceiling"], 5000.0)
        self.assertEqual(kw["n_formants"], 5.0)

    def test_export_receives_defaults(self):
        self._post("/api/review/export", {"dir": "D", "defaults": {
            "pitchFloor": 50, "pitchCeiling": 300, "formantCeiling": 5000,
            "nFormants": 5, "timeStep": 0.005}})
        d = self._seen["export_dir"]["defaults"]
        self.assertEqual(d["pitchFloor"], 50)
        self.assertEqual(d["timeStep"], 0.005)

    def test_export_without_defaults_still_works(self):
        self._post("/api/review/export", {"dir": "D"})
        self.assertIsNone(self._seen["export_dir"]["defaults"])


class MergeNfTest(unittest.TestCase):
    """「多数量择优」的合并/逐音节挑数量逻辑（纯 Python，不依赖 Praat）：

    - 候选 = 4/5/6/7 四套 To Formant (burg) 文本，由 _read_formant_frames 解析后，
      复用 formant_path.choose 的「带宽体检 + 帧间平滑度」打分，逐音节挑数量并合并。
    - 验证点：① 非静音段里默认数量若不过关会被换成更健康的；② 静音段退回默认数量。
    """

    def _write_formant(self, path, nf, per_frame_cands):
        """生成一份极简但合法的 To Formant (burg) 文本（3 帧，每帧 nf 个候选）。"""
        n = len(per_frame_cands)
        dx, x1 = 0.005, 0.0025
        lines = [
            "Formant 1",
            "    xmin = 0.0",
            "    xmax = %.5f" % (x1 + (n - 1) * dx + 0.001),
            "    nx = %d" % n,
            "    dx = %.5f" % dx,
            "    x1 = %.5f" % x1,
            "    maximumNumberOfFormants = %d" % nf,
            "    maximumFrequency = 5500",
        ]
        for fi, cands in enumerate(per_frame_cands, 1):
            lines.append("    frames [%d]:" % fi)
            lines.append("        numberOfFormants = %d" % nf)
            for j, (f, b) in enumerate(cands, 1):
                lines.append("        formant [%d]:" % j)
                lines.append("            frequency = %.6f" % f)
                lines.append("            bandwidth = %.6f" % b)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("\n".join(lines) + "\n")

    def _make_all(self, d, bad_nfs):
        """写齐 4/5/6/7 四套；bad_nfs 里的数量给其中一帧塞一个窄带宽（<30Hz 谐波伪极点）。"""
        key = "x"
        for nf in review_grid.NF_CANDIDATES:
            cands = []
            for fi in range(3):
                fr = []
                for k in range(nf):
                    bw = 20.0 if (nf in bad_nfs and fi == 1) else 60.0 + k * 10
                    fr.append((999.0 - nf * 10 + k + fi, bw))
                cands.append(fr)
            self._write_formant(os.path.join(d, "%s.F%d.formant" % (key, nf)), nf, cands)
        return key

    def test_non_silent_picks_healthy_count(self):
        with tempfile.TemporaryDirectory() as d:
            key = self._make_all(d, bad_nfs={4, 5, 7})   # 只有 6 健康 → 必被选中
            ivs = [{"t0": 0.0, "t1": 0.02, "label": "a"}]  # 非静音
            di = review_grid._closest_index(review_grid.NF_CANDIDATES, 5.0)  # 默认 5
            info = review_grid._merge_nf(
                d, [key], [(None, ivs)], review_grid.NF_CANDIDATES,
                default_index=di, page_names=[key])
            merged = info["forms"][key]
            self.assertEqual(len(merged), 3)
            # 默认(5)与其它 3 套都因窄带宽被淘汰，唯一健康的是 6 → 取 6 的频率
            self.assertAlmostEqual(merged[0]["F1"], 999.0 - 60 + 0, places=2)
            self.assertGreater(info["switched"], 0)
            self.assertGreater(info["intervals"], 0)

    def test_silent_interval_uses_default(self):
        with tempfile.TemporaryDirectory() as d:
            key = self._make_all(d, bad_nfs={4, 6, 7})   # 默认 5 健康
            ivs = [{"t0": 0.0, "t1": 0.02, "label": "sil"}]  # 静音段 → 退回默认
            di = review_grid._closest_index(review_grid.NF_CANDIDATES, 5.0)
            info = review_grid._merge_nf(
                d, [key], [(None, ivs)], review_grid.NF_CANDIDATES,
                default_index=di, page_names=[key])
            merged = info["forms"][key]
            self.assertEqual(len(merged), 3)
            # 默认数量 = 5 → 取 5 的频率（999-50 起）
            self.assertAlmostEqual(merged[0]["F1"], 999.0 - 50 + 0, places=2)
            self.assertEqual(info["switched"], 0)
            self.assertEqual(info["intervals"], 0)


class VowelRefTest(unittest.TestCase):
    """vowel_ref：音段解析、期望共振峰、偏差（比例为主）、性别归一。"""

    def test_parse_segments(self):
        self.assertEqual(vowel_ref.parse_segments("ai"), ["a", "i"])
        self.assertEqual(vowel_ref.parse_segments("iɛ"), ["i", "ɛ"])
        self.assertEqual(vowel_ref.parse_segments("iau"), ["i", "a", "u"])
        self.assertEqual(vowel_ref.parse_segments("ɿ"), ["ɿ"])
        self.assertEqual(vowel_ref.parse_segments("xk"), [])  # 无元音字符

    def test_expected_and_gender(self):
        i_m = vowel_ref.expected("i", "male")
        self.assertIsNotNone(i_m)
        self.assertAlmostEqual(i_m["f1"][0], 342.0)
        i_f = vowel_ref.expected("i", "female")
        self.assertGreater(i_f["f2"][0], i_m["f2"][0])  # 女 F2 更高
        # 无分性别数据 → 回退 general
        self.assertIsNotNone(vowel_ref.expected("ɿ", "male"))
        self.assertIsNone(vowel_ref.expected("zz", "male"))  # 未知元音

    def test_deviation_ratio_dominates(self):
        # 同一元音、不同声道大小：绝对数差很多但比例接近 → 偏差应很小
        ref = vowel_ref.expected("i", "male")  # f1~342 f2~2322
        near = vowel_ref.deviation([400, 2720, 3200], ref)   # 整体放大 ~1.17 倍
        far = vowel_ref.deviation([342, 1100, 3200], ref)     # F2 砍半（错误分析）
        self.assertIsNotNone(near)
        self.assertIsNotNone(far)
        self.assertLess(near, far)
        self.assertLess(near, 0.6)   # 比例接近，偏差小

    def test_gender_of(self):
        self.assertEqual(vowel_ref.gender_of("male"), "male")
        self.assertEqual(vowel_ref.gender_of("女"), "female")
        self.assertEqual(vowel_ref.gender_of(""), None)
        self.assertEqual(vowel_ref.gender_of("M"), "male")


class ChooseRefTest(unittest.TestCase):
    """formant_path.choose 接入「元音参考」与「连续性」两个软判据。"""

    def _frames(self, fs):
        # fs: 每帧 (F1,F2,F3)，带宽统一给 80Hz（健康），生成 choose 要的帧结构
        return [{"t": 0.0025 + i * 0.005, "c": [(f, 80.0) for f in trio]}
                for i, trio in enumerate(fs)]

    def test_ref_picks_matching_vowel(self):
        # 两套候选：A 像 /i/（高 F2），B 像 /a/（低 F2）；参考 = /i/ → 选 A
        a = self._frames([(300, 2300, 2800), (300, 2300, 2800), (300, 2300, 2800)])
        b = self._frames([(700, 1100, 2500), (700, 1100, 2500), (700, 1100, 2500)])
        ref = vowel_ref.expected("i", "male")
        ivs = [(0.0, 0.02, "syl")]   # choose 收的是 partition 产的 (t0,t1,label) 元组
        choices = formant_path.choose([a, b], ivs, default=2, ref=[ref])
        self.assertEqual(choices[0][2], 1)   # 1-based：选第 1 套（/i/ 型）

    def test_continuity_reduces_switch(self):
        # choose 收的是 partition 产的 (t0,t1,label) 元组。
        # 两区间、各 ≥3 帧：区间1 两候选跳变都小→选默认 1；区间2 候选1 跳变大(50)、候选2 跳变小(5)。
        # 无连续性惩罚→切到 2；强惩罚→保持 1（不轻易断裂）。
        mk = lambda fs: [{"t": 0.0025 + i * 0.0025, "c": [(f, 80.0) for f in trio]}
                         for i, trio in enumerate(fs)]
        # 6 帧：前 3 帧∈区间1(0.0–0.01)，后 3 帧∈区间2(0.01–0.02)
        v1 = mk([(300, 1500, 2500), (300, 1500, 2500), (300, 1500, 2500),
                 (300, 1500, 2500), (300, 1500, 2500), (350, 1500, 2500)])
        v2 = mk([(300, 1500, 2500), (300, 1500, 2500), (310, 1500, 2500),
                 (300, 1500, 2500), (300, 1500, 2500), (305, 1500, 2500)])
        ivs = [(0.0, 0.01, "syl"), (0.01, 0.02, "syl")]
        # 区间1：v1 跳变0 < v2 跳变(10/2) → 选 1，prev=1
        c_nopen = formant_path.choose([v1, v2], ivs, default=1, cont_penalty=0.0)
        self.assertEqual(c_nopen[0][2], 1)
        self.assertEqual(c_nopen[1][2], 2)   # 无惩罚，区间2 切到跳变更小的 2
        # 强惩罚：区间2 保持 1（不轻易断裂）
        c_pen = formant_path.choose([v1, v2], ivs, default=1, cont_penalty=60.0)
        self.assertEqual(c_pen[1][2], 1)


class SaveChoiceVowelIpaTest(unittest.TestCase):
    """save_choice 透传 vowel_ipa / speaker，且合并语义不抹掉另一维度。"""

    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.wav = os.path.join(self.d, "x.wav")
        open(self.wav, "w").close()

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def test_vowel_ipa_and_speaker_persist(self):
        r1 = review_choice.save_choice(self.wav, {
            "formant": {"id": "5500-5", "label": "共振峰上限5500Hz·5个"},
            "vowel_ipa": ["a", "i"], "speaker": "female"})
        self.assertTrue(r1["ok"])
        r2 = review_choice.read_choice(self.wav)
        self.assertEqual(r2["choice"]["vowel_ipa"], ["a", "i"])
        self.assertEqual(r2["choice"]["speaker"], "female")

    def test_merge_keeps_vowel_ipa_across_dims(self):
        review_choice.save_choice(self.wav, {"vowel_ipa": ["ə"], "speaker": "male"})
        # 另一维度（pitch）存盘，不带走 vowel_ipa
        review_choice.save_choice(self.wav, {"manual": [{"t": 0.1, "f0": 200}]})
        ch = review_choice.read_choice(self.wav)["choice"]
        self.assertEqual(ch["vowel_ipa"], ["ə"])
        self.assertEqual(ch["speaker"], "male")
        self.assertEqual(ch["manual"], [{"t": 0.1, "f0": 200}])


class SeriesExportTest(unittest.TestCase):
    """第四步导出时间序列：三种重采样方式 + 长表结构。全在 tempfile，不碰真实数据。"""

    def _seed(self, d):
        # 导出靠目录里的 .wav 决定要处理哪些文件（对象文件按主名配对）。
        open(os.path.join(d, "11.wav"), "w").close()
        # 手写一个 .Pitch（3 帧，t=0.0025/0.0075/0.0125）与一个 .Formant（2 帧）。
        with open(_mp(d, "11.Pitch"), "w", encoding="utf-8") as f:
            f.write(
                'File type = "ooTextFile"\nObject class = "Pitch 1"\n'
                "xmin = 0.0 \nxmax = 0.015 \nx1 = 0.0025 \n"
                "nx = 3 \ndx = 0.005 \n"
                "frames []:\n"
                "    frames [1]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 200 \n"
                "    frames [2]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 210 \n"
                "    frames [3]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 220 \n")
        praat_objects.write_formant_text(
            _mp(d, "11.Formant"),
            [{"t": 0.0025, "F1": 500, "F2": 1500, "F3": 2500},
             {"t": 0.0125, "F1": 520, "F2": 1520, "F3": 2520}],
            xmin=0.0, maxn=3)

    def test_raw_exports_all_frames(self):
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0", "F1", "F2", "F3"],
                                             method="raw", out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s,F0,F1,F2,F3")
            # 长表：file, time_s 两固定列 + 4 参数 = 6 列
            self.assertEqual(len(rows[1].split(",")), 6)

    def test_smooth_gives_n_rows_per_file(self):
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], method="smooth", n=5,
                                            out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            # 1 文件标题 + 5 数据行
            self.assertEqual(len(rows), 6)
            # 末行时间应接近 0.0125（t1），证明等间隔铺满
            self.assertAlmostEqual(float(rows[-1].split(",")[1]), 0.0125, places=3)

    def test_interval_samples_by_time(self):
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], method="interval",
                                            interval=0.005, out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            # 0.0025, 0.0075, 0.0125 → 3 行
            self.assertEqual(len(rows), 4)

    def test_no_objects_skips_file(self):
        with tempfile.TemporaryDirectory() as d:
            open(os.path.join(d, "11.wav"), "w").close()
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], out_path=out, ext="wav")
            self.assertFalse(r["ok"])
            self.assertIn("没有可导出", r["error"])

    # ---------------- TextGrid 标注段过滤（2026-09-27 拍板：默认开） ----------------

    TG_TEXT = (
        'File type = "ooTextFile"\nObject class = "TextGrid"\n'
        "xmin = 0 \nxmax = 0.015 \ntiers? size = 1 \nitem []: \n"
        '    item [1]:\n        class = "IntervalTier" \n        name = "syllables" \n'
        "        xmin = 0 \nxmax = 0.015 \nintervals: size = 3 \n"
        "        intervals [1]:\n            xmin = 0 \nxmax = 0.005 \n"
        '            text = "sil" \n'
        "        intervals [2]:\n            xmin = 0.005 \nxmax = 0.01 \n"
        '            text = "a" \n'
        "        intervals [3]:\n            xmin = 0.01 \nxmax = 0.015 \n"
        '            text = "sil" \n')

    def test_tg_filter_default_drops_sil_frames(self):
        """默认开启过滤 + 跳过 sil：3 帧里只剩 t=0.0075（落在标注 a 段内）。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            with open(os.path.join(d, "11.TextGrid"), "w", encoding="utf-8") as f:
                f.write(self.TG_TEXT)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], method="raw",
                                            out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            self.assertEqual(r["tgFiltered"], 1)
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(len(rows), 2)   # 表头 + 1 行
            self.assertAlmostEqual(float(rows[1].split(",")[1]), 0.0075, places=4)

    def test_tg_filter_off_exports_all(self):
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            with open(os.path.join(d, "11.TextGrid"), "w", encoding="utf-8") as f:
                f.write(self.TG_TEXT)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], method="raw",
                                            out_path=out, ext="wav", tg_filter=False)
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(len(rows), 4)   # 表头 + 3 帧

    def test_tg_missing_reports_unfiltered(self):
        """没有 TextGrid 的文件：不过滤、照常导出，并计入 tgMissing。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0"], method="raw",
                                            out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            self.assertEqual(r["tgMissing"], ["11.wav"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(len(rows), 4)

    # ---------------- 双数据源：汇总表（谐波/噪声/能量类） ----------------

    TSV_HEADER = ("file\tlabel\tintervaldID\tt\tf0\tF1\tF2\tF3\t"
                  "H1c\tH2c\tH4c\tA1c\tA2c\tA3c\tH2Ku\tH5Ku\t"
                  "H1u\tH2u\tH4u\tA1u\tA2u\tA3u\t"
                  "H1H2c\tH2H4c\tH1A1c\tH1A2c\tH1A3c\tH2KH5Ku\t"
                  "CPP\tHNR05\tHNR15\tHNR25\tHNR35")
    TSV_COLS = TSV_HEADER.split("\t")

    def _seed_tsv(self, d, wav_name="11.wav", tsv_name="out.tsv"):
        """写一份最小汇总表（3 帧），列名与 initiateTable.praat 的实际表头一致。

        wav_name 决定 file 列写什么（汇总表的 file 列存的是音频文件名）。
        行内容设计成可逐列断言：t=0.0025/0.0075/0.0125，f0=200/210/220，
        谐波差列 H1H2c=4/5/6，HNR05=20.5/21.5/22.5，CPP=15/16/17；
        A1c 用 --undefined-- 验证"未测出→空单元格"。
        """
        def row(t, f0, f1, f2, f3):
            vals = {
                "f0": f0, "F1": f1, "F2": f2, "F3": f3,
                "H1c": 40, "H2c": 35, "H4c": 30, "A1c": "--undefined--", "A2c": 20, "A3c": 18,
                "H2Ku": 12, "H5Ku": 8,
                "H1u": 41, "H2u": 36, "H4u": 31, "A1u": 21, "A2u": 19, "A3u": 17,
                "H1H2c": 4 if t < 0.005 else (5 if t < 0.01 else 6),
                "H2H4c": 5, "H1A1c": 20, "H1A2c": 25, "H1A3c": 22, "H2KH5Ku": 4,
                "CPP": 15 if t < 0.005 else (16 if t < 0.01 else 17),
                "HNR05": 20.5 if t < 0.005 else (21.5 if t < 0.01 else 22.5),
                "HNR15": 30, "HNR25": 25, "HNR35": 22,
            }
            return "\t".join([wav_name, "a", "1", str(t)] +
                             [str(vals.get(c, "")) for c in self.TSV_COLS[4:]])

        lines = [self.TSV_HEADER,
                 row(0.0025, 200, 500, 1500, 2500),
                 row(0.0075, 210, 505, 1505, 2505),
                 row(0.0125, 220, 510, 1510, 2510)]
        with open(os.path.join(d, tsv_name), "w",
                  encoding="utf-8-sig", newline="") as f:
            f.write("\n".join(lines) + "\n")

    def test_tsv_only_source_exports_full_columns(self):
        """只勾汇总表来源的列：没有对象文件也能导出（不再要求 .Pitch/.Formant）。"""
        with tempfile.TemporaryDirectory() as d:
            open(os.path.join(d, "11.wav"), "w").close()
            self._seed_tsv(d, "11.wav")
            out = os.path.join(d, "s.csv")
            cols = ["H1H2c", "HNR05", "CPP", "intensity", "H1c", "A1c"]
            r = series_export.export_series(d, cols, method="raw", out_path=out, ext="wav")
            self.assertTrue(r["ok"], r.get("error"))
            self.assertTrue(r["tsv"].endswith("out.tsv"))
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s," + ",".join(cols))
            c1 = rows[1].split(",")
            self.assertEqual(c1[2], "4")       # H1H2c 第一帧
            self.assertEqual(c1[3], "20.5")    # HNR05
            self.assertEqual(c1[4], "15")      # CPP
            # font: intensity 列在汇总表里没有 → 全空（本夹具没写 intensity）
            self.assertEqual(c1[5], "")        # intensity（夹具未提供）
            self.assertEqual(c1[6], "40")      # H1c
            self.assertEqual(c1[7], "")        # A1c = --undefined-- → 空，不是 0
            self.assertEqual(rows[3].split(",")[2], "6")   # 第三帧 H1H2c

    def test_both_sources_merged_per_column(self):
        """两个源混选：F0 取自 .Pitch、H1H2c 取自汇总表，同一张表的同一行都在。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)              # 11.wav + 11.Pitch(200/210/220) + 11.Formant(F1-F3)
            self._seed_tsv(d, "11.wav")
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["F0", "H1H2c", "HNR05"],
                                            method="raw", out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s,F0,H1H2c,HNR05")
            c1 = rows[1].split(",")
            self.assertEqual(c1[2], "200")     # F0 来自对象
            self.assertEqual(c1[3], "4")       # H1H2c 来自汇总表
            self.assertEqual(c1[4], "20.5")

    def test_tsv_missing_gives_clear_error(self):
        """勾了汇总表来源的列但找不到 out.tsv：只有对象文件时必须明确报错，
        不能静默导出一堆空列。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)              # 只有对象文件，没有 out.tsv
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(d, ["H1H2c"], method="raw",
                                            out_path=out, ext="wav")
            self.assertFalse(r["ok"])
            self.assertIn("汇总表", r["error"])

    def test_column_registry_shape(self):
        """列注册表：47 列、无重复、来源齐全、八组铺平后与 COLUMNS 完全一致。"""
        self.assertEqual(len(series_export.COLUMNS), 47)
        self.assertEqual(len(set(series_export.COLUMNS)), 47)
        flat = [c for _k, _lb, cols in series_export.GROUPS for c in cols]
        self.assertEqual(flat, series_export.COLUMNS)
        for c in series_export.COLUMNS:
            self.assertIn(series_export.col_source(c),
                          ("pitch", "formant", "tsv"), c)
        # 用户示例里的八组顺序与组内列序（照着截图核对）
        self.assertEqual(series_export.GROUPS[0][2], ["F0", "F1", "F2", "F3", "F4", "F5"])
        self.assertEqual(series_export.GROUPS[2][2],
                         ["HNR05", "HNR15", "HNR25", "HNR35", "HNR38", "CPP", "CPPS", "SHR"])
        self.assertEqual(series_export.GROUPS[4][2],
                         ["H1c", "H2c", "H4c", "A1c", "A2c", "A3c", "H2Ku", "H5Ku"])
        self.assertEqual(series_export.GROUPS[5][2], ["B1", "B2", "B3", "B4", "B5"])

    def test_formant_f4_b5_columns(self):
        """2026-09-26 第四步扩展：F1–F5/B1–B5 全列导出。

        顺带锁两件事：write_formant_text 能写 5 阶（平滑存盘不再抹掉 F4/F5）；
        3 阶对象导出时 F4/F5 列是空而不是 0.00（writer 不补零值尾巴）。"""
        with tempfile.TemporaryDirectory() as d:
            open(os.path.join(d, "11.wav"), "w").close()
            praat_objects.write_formant_text(
                _mp(d, "11.Formant"),
                [{"t": 0.0025, "F1": 500, "F2": 1500, "F3": 2500, "F4": 3500, "F5": 4500,
                  "B1": 50, "B2": 60, "B3": 70, "B4": 80, "B5": 90},
                 {"t": 0.0125, "F1": 520, "F2": 1520, "F3": 2520, "F4": 3520, "F5": 4520,
                  "B1": 51, "B2": 61, "B3": 71, "B4": 81, "B5": 91}],
                xmin=0.0, maxn=5)
            out = os.path.join(d, "s.csv")
            cols = ["F0", "F1", "F2", "F3", "F4", "F5", "B1", "B2", "B3", "B4", "B5"]
            r = series_export.export_series(d, cols, method="raw", out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s," + ",".join(cols))
            c1 = rows[1].split(",")
            # F0 没有 .Pitch → 空；F4/F5、B4/B5 都有值
            self.assertEqual(c1[2], "")
            self.assertEqual(c1[6], "3500")
            self.assertEqual(c1[7], "4500")
            self.assertEqual(c1[11], "80")
            self.assertEqual(c1[12], "90")

    def test_formant_3_tiers_export_f4_empty(self):
        """只有 3 阶的对象：F4/F5 列导出为空（writer 不补 0 帧尾巴 → 读回 None）。"""
        with tempfile.TemporaryDirectory() as d:
            open(os.path.join(d, "11.wav"), "w").close()
            praat_objects.write_formant_text(
                _mp(d, "11.Formant"),
                [{"t": 0.0025, "F1": 500, "F2": 1500, "F3": 2500},
                 {"t": 0.0125, "F1": 520, "F2": 1520, "F3": 2520}],
                xmin=0.0, maxn=3)
            out = os.path.join(d, "s.csv")
            r = series_export.export_series(
                d, ["F1", "F2", "F3", "F4", "F5"], method="raw", out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            rows = open(out, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s,F1,F2,F3,F4,F5")
            c1 = rows[1].split(",")
            self.assertEqual(c1[4], "2500")
            self.assertEqual(c1[5], "")   # F4 空
            self.assertEqual(c1[6], "")   # F5 空

    def _seed_second(self, d):
        """第二个音频（12.wav + 12.Pitch，300/310/320 Hz），验证跨文件堆叠。"""
        open(os.path.join(d, "12.wav"), "w").close()
        with open(_mp(d, "12.Pitch"), "w", encoding="utf-8") as f:
            f.write(
                'File type = "ooTextFile"\nObject class = "Pitch 1"\n'
                "xmin = 0.0 \nxmax = 0.015 \nx1 = 0.0025 \n"
                "nx = 3 \ndx = 0.005 \n"
                "frames []:\n"
                "    frames [1]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 300 \n"
                "    frames [2]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 310 \n"
                "    frames [3]:\n        intensity = 1.0 \n        nCandidates = 1 \n"
                "        candidates [1]:\n            frequency = 320 \n")

    def test_layout_per_audio_wide_files(self):
        """每音频一个文件：宽表 time_s × 参数，无 file 列，默认写到「逐音频序列」。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            r = series_export.export_series(d, ["F0", "F1", "F2", "F3"],
                                            method="raw", layout="per_audio", ext="wav")
            self.assertTrue(r["ok"])
            self.assertEqual(r["kind"], "dir")
            self.assertEqual(r["files"], 1)
            self.assertTrue(r["path"].endswith("逐音频序列"))
            p = os.path.join(r["path"], "11.csv")
            rows = open(p, encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "time_s,F0,F1,F2,F3")
            c1 = rows[1].split(",")
            self.assertEqual(len(c1), 5)          # time_s + 4 参数，没有 file 列
            self.assertEqual(c1[1], "200")        # 整数不带 .0
            self.assertEqual(c1[2], "500")
            # 3 帧并集（0.0025/0.0075/0.0125）
            self.assertEqual(len(rows) - 1, 3)

    def test_layout_per_param_stacks_files(self):
        """每参数一个文件：跨音频长表堆叠（11 在前 12 在后），文件数 = 参数数。"""
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            self._seed_second(d)
            r = series_export.export_series(d, ["F0", "F1"], method="raw",
                                            layout="per_param",
                                            out_path=os.path.join(d, "pp"), ext="wav")
            self.assertTrue(r["ok"])
            self.assertEqual(r["kind"], "dir")
            self.assertEqual(r["files"], 2)       # F0.csv + F1.csv
            rows = open(os.path.join(r["path"], "F0.csv"),
                        encoding="utf-8-sig").read().strip().split("\n")
            self.assertEqual(rows[0], "file,time_s,F0")
            self.assertEqual(len(rows) - 1, 6)    # 11 的 3 帧 + 12 的 3 帧
            self.assertEqual(rows[1].split(",")[0], "11")
            self.assertEqual(rows[4].split(",")[0], "12")
            self.assertEqual(rows[4].split(",")[2], "300")

    @unittest.skipUnless(series_export.HAS_OPENPYXL, "需要 openpyxl（未安装则跳过）")
    def test_layout_xlsx_sheets(self):
        """Excel 工作簿：每参数一页 + 说明页（仅装了 openpyxl 的环境验证内容）。"""
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            self._seed(d)
            out = os.path.join(d, "s.xlsx")
            r = series_export.export_series(d, ["F0", "F1"], method="raw",
                                            layout="xlsx", out_path=out, ext="wav")
            self.assertTrue(r["ok"])
            self.assertEqual(r["kind"], "file")
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb.sheetnames[0], "说明")
            self.assertIn("F0", wb.sheetnames)
            self.assertIn("F1", wb.sheetnames)
            sh = wb["F0"]
            self.assertEqual([c.value for c in sh[1]], ["file", "time_s", "F0"])
            self.assertEqual(sh.cell(row=2, column=1).value, "11")
            self.assertEqual(sh.cell(row=2, column=3).value, 200)


class TrackFileNFormantsTest(unittest.TestCase):
    """track_file 的 n_formants 透传（2026-09-26 平滑与候选耦合拍板）：

    详情页「平滑」作用对象跟随勾选的候选——勾了哪条就用它的「上限 + 数量」
    重跑 burg 再 Track，不再写死 5。生成脚本里 ``To Formant (burg)`` 第 2 参
    是 maxNumberOfFormants、第 3 参是 ceiling，直接对脚本文本断言。
    打桩 _run_script（家规：tests/ 内不跑真实 Praat）。
    """

    def setUp(self):
        self._orig = (review_grid._resolve_exe, review_grid._run_script,
                      review_grid.praat_objects.read_formant_text)
        review_grid._resolve_exe = lambda exe=None: "fake-praat.exe"
        self.scripts = []

        def fake_run(exe, out_dir, script, name):
            self.scripts.append(script)
            with open(os.path.join(out_dir, "one.formant"), "w",
                      encoding="utf-8") as f:
                f.write("")
            return None, "OK"

        review_grid._run_script = fake_run
        review_grid.praat_objects.read_formant_text = lambda p: [
            {"t": 0.01, "F1": 500.0, "F2": 1500.0, "F3": 2500.0}]

    def tearDown(self):
        (review_grid._resolve_exe, review_grid._run_script,
         review_grid.praat_objects.read_formant_text) = self._orig

    def _track(self, **kw):
        with tempfile.TemporaryDirectory() as d:
            wav = os.path.join(d, "a.wav")
            open(wav, "w").close()   # 只需存在；Praat 已打桩
            return review_grid.track_file(wav, name="a", **kw)

    def test_n_formants_passed_through(self):
        r = self._track(formant_ceiling=5500.0, n_formants=6)
        self.assertTrue(r["ok"], r.get("error"))
        self.assertIn("To Formant (burg): 0.0, 6, 5500, 0.025, 50",
                      self.scripts[-1])

    def test_default_is_5(self):
        r = self._track(formant_ceiling=5500.0)
        self.assertTrue(r["ok"], r.get("error"))
        self.assertIn("To Formant (burg): 0.0, 5, 5500, 0.025, 50",
                      self.scripts[-1])


if __name__ == "__main__":
    unittest.main(verbosity=2)
